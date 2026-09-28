"""Parse Screener.in company Excel exports → flat fundamentals rows.

Atlas Invest-intel paste accepts CSV/JSON only — not a Screener URL and not a
raw ``.xlsx``. This converter reads the workbook ``Data Sheet`` (raw numbers)
and derives PE / ROE / D/E. FCF is left missing unless an explicit FCF row exists
(never invent Capex).
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

VERSION = "dp.screener_xlsx.v1"


def _last_num(vals: list[Any]) -> float | None:
    nums = [float(v) for v in vals if isinstance(v, (int, float))]
    return nums[-1] if nums else None


def _section_row(
    grid: list[list[Any]],
    label: str,
    *,
    after: str | None = None,
) -> list[Any]:
    started = after is None
    for row in grid:
        if not row or row[0] is None:
            continue
        key = str(row[0]).strip()
        if not started:
            if key == after:
                started = True
            continue
        if key == label:
            return list(row[1:])
    return []


def parse_screener_company_xlsx(
    path: str | Path,
    *,
    symbol: str | None = None,
) -> dict[str, Any]:
    """Extract one fundamentals row from a Screener company workbook.

    PE = Current Price / EPS, EPS = annual Net profit / Adjusted Equity Shares in Cr.
    ROE% = annual Net profit / (Equity Share Capital + Reserves).
    Debt/Equity = Borrowings / equity book.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        return {
            "version": VERSION,
            "ok": False,
            "reason": "openpyxl_missing",
            "error": str(exc),
            "rows": [],
        }

    p = Path(path)
    if not p.is_file():
        return {
            "version": VERSION,
            "ok": False,
            "reason": "file_missing",
            "rows": [],
            "path": str(p),
        }

    wb = load_workbook(p, data_only=True)
    if "Data Sheet" not in wb.sheetnames:
        return {
            "version": VERSION,
            "ok": False,
            "reason": "no_data_sheet",
            "sheets": list(wb.sheetnames),
            "rows": [],
            "honesty": "Need Screener company export with a Data Sheet tab.",
        }

    grid = [list(r) for r in wb["Data Sheet"].iter_rows(values_only=True)]
    company = None
    name_vals = _section_row(grid, "COMPANY NAME")
    if name_vals:
        company = str(name_vals[0]).strip() if name_vals[0] is not None else None

    price = _last_num(_section_row(grid, "Current Price"))
    mcap = _last_num(_section_row(grid, "Market Capitalization"))
    np_ann = _last_num(_section_row(grid, "Net profit", after="PROFIT & LOSS"))
    eq = _last_num(_section_row(grid, "Equity Share Capital", after="BALANCE SHEET")) or 0.0
    res = _last_num(_section_row(grid, "Reserves", after="BALANCE SHEET")) or 0.0
    borr = _last_num(_section_row(grid, "Borrowings", after="BALANCE SHEET"))
    shares = _last_num(_section_row(grid, "No. of Equity Shares", after="BALANCE SHEET"))
    shares_cr = _last_num(_section_row(grid, "Adjusted Equity Shares in Cr"))
    # Explicit FCF rare in this export shape
    fcf = _last_num(_section_row(grid, "Free Cash Flow"))
    if fcf is None:
        fcf = _last_num(_section_row(grid, "FCF"))

    equity = eq + res
    eps = (np_ann / shares_cr) if np_ann and shares_cr else None
    pe = (price / eps) if price and eps else (mcap / np_ann if mcap and np_ann else None)
    roe = (100.0 * np_ann / equity) if np_ann and equity else None
    de = (borr / equity) if borr is not None and equity else None

    sym = (symbol or "").strip().upper()
    if not sym and company:
        # Operator should pass WELCORP.NS; company name alone is not a ticker.
        sym = ""

    row: dict[str, Any] = {
        "symbol": sym or None,
        "company_name": company,
        "source": "screener_export",
        "pe": round(pe, 4) if pe is not None else None,
        "roe": round(roe, 4) if roe is not None else None,
        "debt_to_equity": round(de, 4) if de is not None else None,
        "fcf": fcf,
        "price": price,
        "shares": int(shares) if shares else None,
        "note": (
            "Parsed Screener company xlsx Data Sheet; "
            "PE=CurrentPrice/EPS(annual); FCF left missing unless explicit row"
        ),
    }
    missing = [k for k in ("pe", "fcf", "roe", "debt_to_equity") if row.get(k) is None]
    return {
        "version": VERSION,
        "ok": bool(sym) and pe is not None,
        "path": str(p),
        "company_name": company,
        "symbol": sym or None,
        "rows": [row] if sym else [],
        "derived": row,
        "missing_fields": missing,
        "honesty": (
            "URL paste and raw xlsx are not Invest-intel paste formats. "
            "Use this parser / drop .xlsx into imports/fundamentals/ with a "
            "known symbol, or paste flat CSV. Never invents FCF from CFO+CFI."
        ),
        "reason": None if sym else "symbol_required",
    }


def import_screener_company_xlsx(
    data_dir: str | Path | None,
    path: str | Path,
    *,
    symbol: str,
    program_id: str = "market_intelligence",
    push_to_ira: bool = False,
    research: Any | None = None,
) -> dict[str, Any]:
    """Parse + upsert one Screener company workbook into the fundamentals store."""
    from atlas.investment.fundamentals import SOURCE_SCREENER_EXPORT, upsert_rows

    parsed = parse_screener_company_xlsx(path, symbol=symbol)
    if not parsed.get("rows"):
        return {**parsed, "imported": 0}
    result = upsert_rows(
        data_dir,
        parsed["rows"],
        program_id=program_id,
        source=SOURCE_SCREENER_EXPORT,
        note=str((parsed["rows"][0] or {}).get("note") or "screener xlsx"),
        merge_screener=True,
    )
    ira_out = None
    if push_to_ira and research is not None:
        row = parsed["rows"][0]
        fields = {
            k: row[k]
            for k in ("pe", "fcf", "roe", "debt_to_equity", "price", "shares")
            if row.get(k) is not None
        }
        try:
            ira_out = research.apply_operator_snapshot(
                symbol,
                fields,
                program_id=program_id,
                note=row.get("note") or "screener xlsx",
                evidence_confidence="estimated",
                auto_refresh=True,
            )
        except Exception as exc:  # noqa: BLE001
            ira_out = {"ok": False, "error": type(exc).__name__}
    return {
        **parsed,
        "imported": result.get("imported"),
        "store_path": result.get("path"),
        "ira": ira_out,
    }
